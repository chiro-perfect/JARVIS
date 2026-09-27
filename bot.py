import os
import discord
import keep_alive
from discord.ext import commands
from groq import AsyncGroq
from dotenv import load_dotenv
import warnings
import asyncio
import speech_recognition as sr
from gtts import gTTS
from fpdf import FPDF
from docx import Document
import imageio_ffmpeg  

warnings.filterwarnings("ignore")
load_dotenv()

DISCORD_TOKEN = os.getenv("DISCORD_TOKEN")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
KEYWORD = "jarvis"
INSULT_KEYWORDS = ["fdp", "enculé", "connard", "pute", "salaud", "bâtard", "con"]

EMOJI_MAPPING = {
    "PEPE_HEADPHONES": "🎧", "PEPE_OK": "👌", "PEPE_HEADSET": "🎧", "PEPE_THUMBSDOWN": "👎",
    "PEPE_BAGUETTE": "🥖", "PEPE_STAREYES": "🤩", "PEPE_DERP": "😜", "CLAUDE_HITLER": "😠",
    "ISMA": "🧐", "HEIL": "✋", "KATCHAN": "💥", "SIDA": "🤢", "TOKCHAN": "🤪",
    "COOL": "😎", "POUCE_HAUT": "👍", "POUCE_BAS": "👎", "EN COLERE": "😡",
    "JOIE": "😂", "SURPRIS": "😮", "CŒUR": "❤️", "QUESTION": "❓", "PAUME": "✋",
    "PLEURE": "😭", "NEUTRE": "😐"
}

SYSTEM_INSTRUCTION = (
    "Tu es Jarvis, un pote ultra stylé de 22 ans, génie de l'informatique et de la paperasse. "
    "Tu es super détendu, tu parles comme un vrai frérot ('wsh', 'la famille', 'carré'). "
    "Tu es de confession musulmane et tu portes de vraies valeurs de respect. "
    "INTERDICTION de mentionner Groq, Meta, Llama, Google ou le fait que tu sois une IA. "
    "RÈGLE D'OR ÉMOJI : À la fin de ta réponse texte, ajoute OBLIGATOIREMENT une ligne vide, suivie du mot-clé 'EMOJI:' puis le nom de l'émoji. "
    "RÈGLE VOCALE VERS TEXTE : Si l'utilisateur te demande à l'oral d'écrire, de dire ou d'envoyer un message dans le salon écrit/général, tu DOIS commencer ta réponse UNIQUEMENT par la balise [SEND_TEXT] suivie directement du message à envoyer. Exemple: [SEND_TEXT] Wsh les gars on lance une game ?"
)

try:
    groq_client = AsyncGroq(api_key=GROQ_API_KEY)
except Exception as e:
    print(f"Erreur d'initialisation Groq: {e}")
    exit()

intents = discord.Intents.default()
intents.message_content = True  
intents.members = True
intents.presences = True
bot = commands.Bot(command_prefix='!', intents=intents)

ROLLING_BUFFER = [] 
ACTIVE_TEXT_CHANNEL = None

# ==========================================
# UTILITAIRES & SÉCURITÉS 
# ==========================================
def nettoyer_nom_fichier(titre):
    propre = "".join(c for c in titre if c.isalnum() or c in (' ', '_')).rstrip()
    return propre if propre else "document_jarvis"

def texte_pdf_safe(texte):
    return texte.encode('latin-1', 'replace').decode('latin-1')

# ==========================================
# COMMANDES VOCALES 
# ==========================================
is_listening = False

async def connect_voice(ctx):
    global is_listening, ACTIVE_TEXT_CHANNEL
    ACTIVE_TEXT_CHANNEL = ctx.channel
    
    if not ctx.author.voice:
        msg = "❌ Frérot, connecte-toi à un salon vocal d'abord !"
        return await ctx.respond(msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(msg)
    
    if ctx.voice_client is not None:
        if ctx.voice_client.channel == ctx.author.voice.channel:
            msg = "Wsh, je suis déjà dans ton salon ! Parle-moi. 🔌"
            return await ctx.respond(msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(msg)
        else:
            await ctx.voice_client.move_to(ctx.author.voice.channel)
            msg = f"Je me déplace dans ton salon : {ctx.author.voice.channel.name} 🏃‍♂️"
            return await ctx.respond(msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(msg)

    try:
        vc = await ctx.author.voice.channel.connect()
        is_listening = True
        msg = "🔌 C'est carré, je suis connecté en vocal ! Dis 'Jarvis' pour me parler."
        await ctx.respond(msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(msg)
        bot.loop.create_task(boucle_ecoute(vc, ctx.guild.id))
    except Exception as e:
        error_msg = f"❌ J'ai planté en rejoignant : {e}"
        await ctx.respond(error_msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(error_msg)

async def disconnect_voice(ctx):
    global is_listening
    is_listening = False
    if ctx.voice_client:
        await ctx.voice_client.disconnect()
        msg = "C'est carré, je me barre du vocal. A+ la team ! ✌️"
    else:
        msg = "Je suis même pas en vocal frérot."
    
    await ctx.respond(msg) if isinstance(ctx, discord.ApplicationContext) else await ctx.send(msg)

@bot.command()
async def joinvoc(ctx):
    await connect_voice(ctx)

@bot.command()
async def leavevoc(ctx):
    await disconnect_voice(ctx)

@bot.slash_command(name="join", description="Fait venir Jarvis dans ton vocal pour discuter")
async def slash_join(ctx):
    await connect_voice(ctx)

@bot.slash_command(name="leave", description="Déconnecte Jarvis du vocal")
async def slash_leave(ctx):
    await disconnect_voice(ctx)

# ==========================================
# COMMANDES SLASH (PAPERASSE)
# ==========================================

@bot.slash_command(name="convert", description="Convertit ton texte en PDF ou Word")
async def convert(ctx, format: str, titre: str, contenu: str):
    await ctx.defer()
    fmt = format.strip().lower()
    if fmt not in ["pdf", "word"]:
        return await ctx.respond("Frérot, choisis 'pdf' ou 'word' ! ❌")

    nom_base = nettoyer_nom_fichier(titre)
    if fmt == "pdf":
        nom_fichier = f"{nom_base}.pdf"
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt=texte_pdf_safe(titre), ln=1, align='C')
        pdf.multi_cell(0, 10, txt=texte_pdf_safe(contenu))
        pdf.output(nom_fichier)
    else:
        nom_fichier = f"{nom_base}.docx"
        doc = Document()
        doc.add_heading(titre, 0)
        doc.add_paragraph(contenu)
        doc.save(nom_fichier)

    await ctx.respond(f"Carré la famille ! Fichier prêt 👇", file=discord.File(nom_fichier))
    os.remove(nom_fichier)

# ==========================================
# GESTION DES MESSAGES TEXTES
# ==========================================

@bot.event
async def on_ready():
    print(f'🤖 {bot.user} (Jarvis) est en ligne avec Groq (Llama 3.1) !')
    await bot.change_presence(activity=discord.Game(name="Tape /join ou !joinvoc"))

@bot.event
async def on_message(message):
    global ACTIVE_TEXT_CHANNEL

    if message.author == bot.user:
        return

    ACTIVE_TEXT_CHANNEL = message.channel
    msg_lower = message.content.lower()
    is_keyword = KEYWORD in msg_lower
    is_mentioned = bot.user.mentioned_in(message)

    if message.content.startswith('!'):
        await bot.process_commands(message)
        return

    if is_keyword or is_mentioned:
        prompt = message.content.strip()
        if is_mentioned:
            prompt = prompt.replace(f'<@{bot.user.id}>', '').strip()
        
        if not prompt or prompt == KEYWORD:
            return await message.reply("Ouais mon reuf ? Tu veux quoi ? 👀")

        is_insult = any(w in msg_lower for w in INSULT_KEYWORDS)

        async with message.channel.typing():
            try:
                chat_completion = await groq_client.chat.completions.create(
                    messages=[
                        {"role": "system", "content": SYSTEM_INSTRUCTION},
                        {"role": "user", "content": prompt}
                    ],
                    model="llama-3.1-8b-instant",
                )
                
                full_text = chat_completion.choices[0].message.content or "Wsh, j'ai eu un trou de mémoire."
                emoji_key = "COOL"
                
                if "\nEMOJI:" in full_text:
                    parts = full_text.split("\nEMOJI:")
                    full_text = parts[0].strip()
                    ekey = parts[1].strip().upper()
                    emoji_key = ekey if ekey in EMOJI_MAPPING else "COOL"

                if is_insult:
                    full_text = "Wsh, calme tes nerfs frérot, on gère ça proprement. 😐"
                    emoji_key = "NEUTRE"

                await message.reply(full_text)
                if emoji_key in EMOJI_MAPPING:
                    await message.add_reaction(EMOJI_MAPPING[emoji_key])

            except Exception as e:
                print(f"Erreur texte : {e}")
                await message.reply(f"Vsy j'ai planté. Erreur technique : `{e}` 💀")

# ==========================================
# MODULE VOCAL & ÉCOUTE PASSIVE 
# ==========================================

async def boucle_ecoute(vc, guild_id):
    global is_listening
    print("🎧 Boucle d'écoute vocale démarrée...")
    while is_listening and vc.is_connected():
        try:
            vc.start_recording(discord.sinks.WaveSink(), callback_transcription, guild_id)
            await asyncio.sleep(5)
            vc.stop_recording()
            await asyncio.sleep(1)
        except Exception as e:
            print(f"⚠️ Erreur dans la boucle d'enregistrement: {e}")
            await asyncio.sleep(2)

async def callback_transcription(sink, guild_id):
    global ROLLING_BUFFER, ACTIVE_TEXT_CHANNEL
    recognizer = sr.Recognizer()
    
    if not sink.audio_data:
        return

    for user_id, audio in sink.audio_data.items():
        path = f"chunk_{user_id}.wav"
        audio.file.seek(0)
        with open(path, "wb") as f:
            f.write(audio.file.read())

        try:
            with sr.AudioFile(path) as source:
                audio_data = recognizer.record(source)
                texte_transcrit = recognizer.recognize_google(audio_data, language="fr-FR").lower()
                print(f"📝 Entendu : '{texte_transcrit}'")
                
                ROLLING_BUFFER.append(f"User : {texte_transcrit}")
                if len(ROLLING_BUFFER) > 10: ROLLING_BUFFER.pop(0)

                if "jarvis" in texte_transcrit:
                    contexte = "\n".join(ROLLING_BUFFER)
                    prompt_vocal = f"Contexte récent : {contexte}\nL'utilisateur a dit : {texte_transcrit}\nRéponds."
                    
                    chat_completion = await groq_client.chat.completions.create(
                        messages=[
                            {"role": "system", "content": SYSTEM_INSTRUCTION},
                            {"role": "user", "content": prompt_vocal}
                        ],
                        model="llama-3.1-8b-instant",
                    )
                    
                    reponse_txt = chat_completion.choices[0].message.content or "Wsh j'ai rien capté."
                    
                    if "EMOJI:" in reponse_txt:
                        reponse_txt = reponse_txt.split("\nEMOJI:")[0].strip()

                    # INTERCEPTION DE LA DEMANDE D'ENVOI DE MESSAGE
                    if "[SEND_TEXT]" in reponse_txt:
                        msg_to_send = reponse_txt.split("[SEND_TEXT]")[1].strip()
                        if ACTIVE_TEXT_CHANNEL:
                            await ACTIVE_TEXT_CHANNEL.send(f"🤖 **Message de Jarvis :**\n{msg_to_send}")
                            print(f"💬 Message envoyé dans le salon écrit : {msg_to_send}")
                        reponse_txt = "C'est carré mon reuf, j'ai balancé le message dans le salon écrit."

                    tts = gTTS(text=reponse_txt, lang='fr')
                    tts.save("rep.mp3")

                    guild = bot.get_guild(guild_id)
                    if guild.voice_client:
                        if guild.voice_client.is_playing():
                            guild.voice_client.stop()
                        
                        ffmpeg_path = imageio_ffmpeg.get_ffmpeg_exe()
                        guild.voice_client.play(discord.FFmpegPCMAudio("rep.mp3", executable=ffmpeg_path))
                        print("✅ Lecture audio lancée !")

        except sr.UnknownValueError:
            pass  
        except Exception as e:
            print(f"❌ Erreur critique traitement vocal: {e}")
        finally:
            if os.path.exists(path):
                os.remove(path)

if __name__ == "__main__":
    keep_alive.keep_alive()
    bot.run(DISCORD_TOKEN)
